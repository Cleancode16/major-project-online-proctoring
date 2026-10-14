"""
Word (.docx) version of a proctoring report (flow.txt §C.4 "Word report", DOCX1). Uses python-docx.
Same sections as the PDF (aggregation/pdf_report.py); voice clips are left out (audio cannot be played in a
document), and the charts are given as tables.

  build_docx(report, report_dir, out_path)  -> out_path
  CLI:  python -m aggregation.docx_report <report.json> [out.docx] [--evidence-dir DIR]
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from aggregation.pdf_report import (DOC_GROUPS, GROUPS, VERDICT_STYLE, _find_evidence_dir, counter_cells, fmt_t,
                                    limits_text, nice, points_text, score_rows)

INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
VERDICT_HEX = {"VALID": ("DCFCE7", RGBColor(0x16, 0x65, 0x34)),
               "SUSPICIOUS": ("FEF3C7", RGBColor(0x92, 0x40, 0x0E)),
               "INVALID": ("FEE2E2", RGBColor(0x99, 0x1B, 0x1B))}
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def txt(v):
    """Text python-docx accepts (XML forbids control characters)."""
    return _CTRL.sub("", str(v if v is not None else ""))


# ------------------------------------------------------------------------------------- helpers
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def para(doc_or_cell, text="", size=9, bold=False, color=INK, align=None, space_after=2):
    p = doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(txt(text))
        r.font.size, r.bold, r.font.color.rgb = Pt(size), bold, color
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, level=2):
    h = doc.add_heading(txt(text), level=level)
    for r in h.runs:
        r.font.color.rgb = INK
    return h


def table(doc, rows, widths_mm=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            r = cell.paragraphs[0].add_run(txt(val))
            r.font.size = Pt(8)
            if header and i == 0:
                r.bold = True
                shade(cell, "F1F5F9")
            if widths_mm:
                cell.width = Mm(widths_mm[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ------------------------------------------------------------------------------------- sections
def _header(doc, report):
    m = report.get("meta", {})
    heading(doc, f"Proctoring report - candidate {m.get('candidate', '-')}", level=1)
    para(doc, f"Session {m.get('session_id', '-')}  |  {m.get('created', '')}  |  mode {m.get('mode', 'recorded')}"
              f"  |  profile {m.get('profile', '')}", size=9, color=MUTED, space_after=8)


def _verdict(doc, report):
    v = report["verdict"]
    fill, fg = VERDICT_HEX.get(v["label"], VERDICT_HEX["SUSPICIOUS"])
    word = VERDICT_STYLE.get(v["label"], VERDICT_STYLE["SUSPICIOUS"])[2]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    left, right = t.cell(0, 0), t.cell(0, 1)
    left.width, right.width = Mm(140), Mm(40)
    for c in (left, right):
        shade(c, fill)
        c.text = ""
    p = left.paragraphs[0]
    r = p.add_run(txt(v["label"]))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), fg
    r = p.add_run(f"  ({word})")
    r.font.size, r.font.color.rgb = Pt(9), fg
    para(left, v["headline"], size=9, bold=True, color=fg)
    for reason in v["reasons"][1:]:
        para(left, f"- {reason}", size=8, color=fg)
    if v.get("score_explanation"):
        para(left, v["score_explanation"], size=7.5, color=fg)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Violation score")
    r.font.size, r.font.color.rgb = Pt(8), fg
    para(right, f"{v['score']:.1f}", size=22, bold=True, color=fg, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    para(right, "of 100", size=7, color=fg, align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _summary(doc, report):
    vid, aud = report.get("video") or {}, report.get("audio") or {}
    b = vid.get("baseline") or {}
    audio = (aud.get("status") or "not analysed").replace("_", " ")
    if aud.get("status") == "ok":
        audio += f", {aud.get('duration_s', 0):.0f} s, speech {aud.get('speech_s', 0):.1f} s"
    events = str(len(report.get("timeline", [])))
    if b:
        events += f"  (baseline yaw {b['yaw']:.1f}, pitch {b['pitch']:.1f}, roll {b['roll']:.1f} deg)"
    table(doc, [["Video analysed", "Audio", "Events"],
                [f"{vid.get('duration_s', 0):.0f} s, {vid.get('n_samples', 0)} samples @ "
                 f"{(vid.get('sample_rate_hz') or 0):.2f}/s", audio, events]])
    for w in report.get("warnings", []):
        para(doc, f"Warning: {w}", size=9, color=RGBColor(0x92, 0x40, 0x0E))


def _scoring(doc, report):
    rows = score_rows(report)
    if not rows:
        return
    v = report["verdict"]
    heading(doc, "Scoring")
    p = doc.add_paragraph()
    r = p.add_run(txt(f"Overall score {v['score']:.1f} / 100 - {v['label']}. "))
    r.bold, r.font.size = True, Pt(9)
    r = p.add_run(txt(v.get("score_explanation", "")))
    r.font.size = Pt(9)
    trs = [["Violation", "Severity", "Points", "Explanation"]]
    for t, sev, pts, mx, expl in rows:
        trs.append([nice(t), sev, f"{pts:.1f} / {mx}" if pts is not None else f"- / {mx}", expl])
    table(doc, trs, [34, 16, 18, 112])


def _evidence_counters(doc, report):
    counters = report.get("evidence_counters") or {}
    if not counters:
        return
    sv = report.get("session_violation") or {}
    up = report.get("evidence_upload") or {}
    heading(doc, "Session violation and evidence counters")
    if sv.get("violated"):
        para(doc, "SESSION MARKED AS VIOLATED - limits exceeded for: "
                  + ", ".join(nice(g) for g in sv.get("groups_exceeded", [])), bold=True, color=RGBColor(0x99, 0x1B, 0x1B))
    rows = [["Group", "Events", "% of session", "Limit", "Exceeded", "Upload", "Target / reason"]]
    for g, label in GROUPS:
        c, u = counters.get(g, {}), up.get(g, {})
        pct, lim = counter_cells(c)
        rows.append([label, str(c.get("count", 0)), pct, lim, "YES" if c.get("exceeded") else "no",
                     u.get("status", "-"), u.get("error") or u.get("target") or u.get("reason") or ""])
    table(doc, rows, [32, 13, 18, 26, 16, 17, 58])


def _per_violation(doc, report):
    heading(doc, "Violations by type")
    rows = [["Violation", "Source", "%", "Numerator / denominator", "Seconds", "Events", "Severity", "Points"]]
    for t, s in report["per_violation"].items():
        rows.append([nice(t), s.get("source", ""),
                     f"{s['percentage']:.1f}" if s.get("measurable") else "n/a",
                     f"{s['numerator']} / {s['denominator']} {s['unit']}" if s.get("measurable") else "not measurable",
                     f"{s.get('duration_s', 0):.1f}", str(s.get("count", 0)),
                     s.get("severity") or ("critical" if s.get("critical") else ""), points_text(s)])
    table(doc, rows, [36, 13, 12, 46, 16, 14, 14, 29])


def _evidence(doc, report, evidence_dir):
    ev = report.get("evidence") or {}
    if not ev:
        return
    heading(doc, "Evidence")
    for g, label in DOC_GROUPS:                     # voice clips are not part of the document (audio only)
        items = (ev.get(g) or {}).get("items", [])
        heading(doc, f"{label} ({len(items)} event(s))", level=3)
        if not items:
            para(doc, "No events.", size=8, color=MUTED)
            continue
        n_rows = (len(items) + 2) // 3
        t = doc.add_table(rows=n_rows, cols=3)
        for k, it in enumerate(items):
            cell = t.cell(k // 3, k % 3)
            cell.text = ""
            path = evidence_dir / it["file"] if it.get("file") else None
            added = False
            if path is not None and path.is_file():
                try:
                    cell.paragraphs[0].add_run().add_picture(str(path), width=Mm(55))
                    added = True
                except Exception:
                    pass
            if not added:
                para(cell, it.get("note") or "image not available", size=7.5, color=MUTED)
            para(cell, _caption(g, it), size=7.5, color=MUTED, space_after=6)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _caption(group, it):
    base = f"{fmt_t(it['start_s'])} - {fmt_t(it['end_s'])} ({it['duration_s']:.1f} s)"
    if group == "object":
        objs = ", ".join(f"{o['label']} {100 * o['confidence']:.0f}%" for o in it.get("objects", []))
        return f"{base}\n{objs or 'object'}"
    if group == "multiple_faces":
        return (f"{base}\n{it.get('face_count', '?')} face(s), {it.get('person_count', '?')} person(s), "
                f"max {it.get('max_persons', '?')}")
    if group == "head_pose":
        a, lim = it.get("angles_deg", {}), it.get("allowed_deg", {})
        return (f"{base}\nyaw {a.get('yaw', '-')}, pitch {a.get('pitch', '-')}, roll {a.get('roll', '-')} deg "
                f"(limits {lim.get('yaw')}/{lim.get('pitch')}/{lim.get('roll')}); allowed {it.get('allowed_duration_s')} s")
    return base


def _events(doc, report):
    tl = report.get("timeline", [])
    if not tl:
        return
    heading(doc, "All events")
    rows = [["Start", "End", "Duration", "Source", "Violation", "Details"]]
    for e in tl:
        rows.append([fmt_t(e["start_s"]), fmt_t(e["end_s"]), f"{e['duration_s']:.1f} s", e["source"], nice(e["type"]),
                     json.dumps(e.get("details") or {})[:160]])
    table(doc, rows, [18, 18, 18, 14, 38, 74])


def _thresholds(doc, report):
    th = report.get("thresholds") or {}
    hp, gz, ob, ev = th.get("head_pose", {}), th.get("gaze", {}), th.get("objects", {}), th.get("evidence", {})
    heading(doc, "Thresholds used")
    lines = [
        f"Head pose: yaw {hp.get('yaw_thr')} / pitch {hp.get('pitch_thr')} / roll {hp.get('roll_thr')} deg, "
        f"for {hp.get('persist_s')} s (method {hp.get('method', 'pnp')})",
        f"Gaze: {gz.get('h_thr_deg')} deg horizontal / {gz.get('v_thr_deg')} deg vertical, for {gz.get('persist_s')} s",
        f"Face match threshold (cosine): {th.get('face', {}).get('cos_threshold')}; "
        f"voice match threshold: {th.get('speaker', {}).get('verify_thr')}",
        "Objects: " + ", ".join(f"{name} >= {ob.get('conf', {}).get(cls, ob.get('conf', {}).get(str(cls), '?'))}"
                                for cls, name in (ob.get("classes") or {}).items()),
        "Session violated when: " + limits_text(ev),
    ]
    for s in lines:
        para(doc, f"- {s}", size=8, color=MUTED)


def _footer(doc, report):
    m = report.get("meta", {})
    p = doc.sections[0].footer.paragraphs[0]
    r = p.add_run(txt(f"Online Proctoring - candidate {m.get('candidate', '-')} - session {m.get('session_id', '-')}"))
    r.font.size, r.font.color.rgb = Pt(7), MUTED


# ------------------------------------------------------------------------------------- build
def build_docx(report, evidence_dir, out_path):
    """Write the Word file for `report` (dict). evidence_dir = folder containing evidence/<group>/... files."""
    out_path = Path(out_path)
    evidence_dir = Path(evidence_dir)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(16)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)
    doc.core_properties.title = txt(f"Proctoring report {report.get('meta', {}).get('candidate', '')}")
    doc.core_properties.author = "Online Proctoring System"
    _header(doc, report)
    _verdict(doc, report)
    _summary(doc, report)
    _scoring(doc, report)
    _evidence_counters(doc, report)
    _per_violation(doc, report)
    _evidence(doc, report, evidence_dir)
    _events(doc, report)
    _thresholds(doc, report)
    _footer(doc, report)
    doc.save(str(out_path))
    return out_path


def write_report_docx(report_json_path):
    """Build report.docx next to report.json (used by the API)."""
    p = Path(report_json_path)
    report = json.loads(p.read_text(encoding="utf-8"))
    return build_docx(report, p.parent, p.with_suffix(".docx"))


def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    src = Path(argv[0])
    out = Path(argv[1]) if len(argv) > 1 and not argv[1].startswith("--") else src.with_suffix(".docx")
    report = json.loads(src.read_text(encoding="utf-8"))
    ev_dir = Path(argv[argv.index("--evidence-dir") + 1]) if "--evidence-dir" in argv else _find_evidence_dir(src, report)
    build_docx(report, ev_dir, out)
    print(f"Word file written: {out} (evidence from {ev_dir})")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
